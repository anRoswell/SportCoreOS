import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(11, 15, 25)
bg.line.fill.solid()
bg.line.color.rgb = RGBColor(11, 15, 25)

# Test card
card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(4))
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(17, 24, 39)
card.line.fill.solid()
card.line.color.rgb = RGBColor(16, 185, 129)

# Test image
img_path = "/Users/sectic/Documents/Codigos Fuentes/SportCoreOS/assets/screenshots/01_inicio_dashboard_raw.png"
im = Image.open(img_path)
w, h = im.size
target_h = Inches(4.2)
target_w = Inches(4.2 * (w / h))
target_left = Inches(7.8) + (Inches(4.733) - target_w) / 2
target_top = Inches(2.25) + (Inches(4.7) - target_h) / 2
slide.shapes.add_picture(img_path, target_left, target_top, width=target_w, height=target_h)

prs.save("/tmp/test_clean.pptx")
print("Saved cleanly!")
